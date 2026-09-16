from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

from backend.app.schemas.generated_outputs import (
    PresentationContent,
    PresentationSlide,
)
from backend.app.services.rendering.base import RenderedResult, Renderer


class PresentationRenderer(Renderer):
    """
    Production PowerPoint renderer for PRISM.

    Responsibilities:
    - Convert verified PresentationContent into an editable PPTX.
    - Provide consistent professional visual hierarchy.
    - Select layouts dynamically based on slide content.
    - Render key statistics without inventing information.
    - Preserve speaker notes and source references.
    - Never generate or modify factual content.
    """

    output_type = "presentation"

    # ---------------------------------------------------------
    # Presentation dimensions
    # ---------------------------------------------------------

    SLIDE_WIDTH = Inches(13.333)
    SLIDE_HEIGHT = Inches(7.5)

    # ---------------------------------------------------------
    # Theme
    # ---------------------------------------------------------

    NAVY = RGBColor(15, 23, 42)
    BLUE = RGBColor(37, 99, 235)
    BLUE_LIGHT = RGBColor(239, 246, 255)

    WHITE = RGBColor(255, 255, 255)
    BLACK = RGBColor(15, 23, 42)

    GRAY_700 = RGBColor(51, 65, 85)
    GRAY_600 = RGBColor(71, 85, 105)
    GRAY_500 = RGBColor(100, 116, 139)
    GRAY_300 = RGBColor(203, 213, 225)
    GRAY_200 = RGBColor(226, 232, 240)
    GRAY_100 = RGBColor(241, 245, 249)

    GREEN = RGBColor(22, 163, 74)
    GREEN_LIGHT = RGBColor(240, 253, 244)

    AMBER = RGBColor(217, 119, 6)
    AMBER_LIGHT = RGBColor(255, 251, 235)

    RED = RGBColor(220, 38, 38)
    RED_LIGHT = RGBColor(254, 242, 242)

    FONT = "Aptos"

    # ---------------------------------------------------------
    # Main renderer
    # ---------------------------------------------------------

    def render(
        self,
        content: PresentationContent,
        output_dir: Path,
    ) -> RenderedResult:

        output_dir.mkdir(parents=True, exist_ok=True)

        presentation = Presentation()

        presentation.slide_width = self.SLIDE_WIDTH
        presentation.slide_height = self.SLIDE_HEIGHT

        self._remove_default_slide(presentation)

        # -----------------------------------------------------
        # Render generated slides
        # -----------------------------------------------------

        for index, slide_data in enumerate(content.slides):

            if index == 0:
                self._add_title_slide(
                    presentation,
                    content,
                    slide_data,
                )
                continue

            self._add_dynamic_content_slide(
                presentation,
                slide_data,
            )

        # -----------------------------------------------------
        # Save
        # -----------------------------------------------------

        file_path = output_dir / "presentation.pptx"

        presentation.save(file_path)

        return RenderedResult(
            output_type=self.output_type,
            file_path=file_path,
            mime_type=(
                "application/vnd.openxmlformats-officedocument."
                "presentationml.presentation"
            ),
            metadata={
                "slide_count": len(presentation.slides),
                "generated_slide_count": len(content.slides),
                "format": "pptx",
                "editable": True,
            },
        )

    # =========================================================
    # Presentation setup
    # =========================================================

    @staticmethod
    def _remove_default_slide(
        presentation: Presentation,
    ) -> None:

        if not presentation.slides:
            return

        slide_id = presentation.slides._sldIdLst[0].rId

        presentation.part.drop_rel(slide_id)

        del presentation.slides._sldIdLst[0]

    # =========================================================
    # TITLE SLIDE
    # =========================================================

    def _add_title_slide(
        self,
        presentation: Presentation,
        content: PresentationContent,
        slide_data: PresentationSlide,
    ) -> None:

        slide = presentation.slides.add_slide(
            presentation.slide_layouts[6]
        )

        background = slide.background.fill
        background.solid()
        background.fore_color.rgb = self.WHITE

        # Left visual panel
        panel = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            0,
            0,
            Inches(0.18),
            self.SLIDE_HEIGHT,
        )

        panel.fill.solid()
        panel.fill.fore_color.rgb = self.BLUE
        panel.line.fill.background()

        # PRISM label
        label = slide.shapes.add_textbox(
            Inches(0.85),
            Inches(0.75),
            Inches(4.0),
            Inches(0.4),
        )

        self._set_text(
            label.text_frame,
            "PRISM",
            font_size=15,
            bold=True,
            color=self.BLUE,
        )

        # Title
        title = slide.shapes.add_textbox(
            Inches(0.85),
            Inches(1.55),
            Inches(10.8),
            Inches(1.7),
        )

        self._set_text(
            title.text_frame,
            content.title,
            font_size=34,
            bold=True,
            color=self.NAVY,
        )

        # Subtitle
        subtitle = content.subtitle or slide_data.subtitle

        if subtitle:
            subtitle_box = slide.shapes.add_textbox(
                Inches(0.88),
                Inches(3.35),
                Inches(9.8),
                Inches(0.9),
            )

            self._set_text(
                subtitle_box.text_frame,
                subtitle,
                font_size=18,
                color=self.GRAY_600,
            )

        # Audience
        if content.audience:

            audience_box = slide.shapes.add_textbox(
                Inches(0.88),
                Inches(5.45),
                Inches(8.0),
                Inches(0.45),
            )

            self._set_text(
                audience_box.text_frame,
                f"Audience  •  {content.audience}",
                font_size=11,
                color=self.GRAY_500,
            )

        # Decorative geometry
        circle = slide.shapes.add_shape(
            MSO_SHAPE.OVAL,
            Inches(10.75),
            Inches(1.35),
            Inches(1.65),
            Inches(1.65),
        )

        circle.fill.solid()
        circle.fill.fore_color.rgb = self.BLUE_LIGHT
        circle.line.fill.background()

        circle_2 = slide.shapes.add_shape(
            MSO_SHAPE.OVAL,
            Inches(11.65),
            Inches(2.25),
            Inches(0.75),
            Inches(0.75),
        )

        circle_2.fill.solid()
        circle_2.fill.fore_color.rgb = self.BLUE
        circle_2.line.fill.background()

        # Opening slide content
        if slide_data.content:

            body = slide.shapes.add_textbox(
                Inches(0.88),
                Inches(4.15),
                Inches(9.0),
                Inches(1.0),
            )

            self._set_text(
                body.text_frame,
                slide_data.content[0],
                font_size=15,
                color=self.GRAY_700,
            )

        self._add_slide_number(
            slide,
            slide_data.slide_number,
        )

        self._add_notes(
            slide,
            slide_data.speaker_notes,
        )

    # =========================================================
    # DYNAMIC CONTENT SLIDES
    # =========================================================

    def _add_dynamic_content_slide(
        self,
        presentation: Presentation,
        slide_data: PresentationSlide,
    ) -> None:

        layout = self._select_layout(slide_data)

        if layout == "stat":
            self._add_stat_slide(
                presentation,
                slide_data,
            )

        elif layout == "timeline":
            self._add_timeline_slide(
                presentation,
                slide_data,
            )

        elif layout == "two_column":
            self._add_two_column_slide(
                presentation,
                slide_data,
            )

        else:
            self._add_standard_slide(
                presentation,
                slide_data,
            )

    # =========================================================
    # Layout selection
    # =========================================================

    def _select_layout(
        self,
        slide_data: PresentationSlide,
    ) -> str:

        title = slide_data.title.lower()

        timeline_words = (
            "timeline",
            "sequence",
            "process",
            "milestone",
            "chronology",
        )

        stat_words = (
            "impact",
            "metrics",
            "statistics",
            "key facts",
            "findings",
            "scope",
        )

        if any(word in title for word in timeline_words):
            return "timeline"

        if slide_data.key_stat:
            return "stat"

        if any(word in title for word in stat_words):
            return "two_column"

        if len(slide_data.content) >= 4:
            return "two_column"

        return "standard"

    # =========================================================
    # STANDARD SLIDE
    # =========================================================

    def _add_standard_slide(
        self,
        presentation: Presentation,
        slide_data: PresentationSlide,
    ) -> None:

        slide = self._create_base_slide(
            presentation,
            slide_data,
        )

        content_box = slide.shapes.add_textbox(
            Inches(0.85),
            Inches(1.85),
            Inches(11.5),
            Inches(4.65),
        )

        frame = content_box.text_frame
        frame.clear()
        frame.word_wrap = True

        for index, point in enumerate(slide_data.content):

            paragraph = (
                frame.paragraphs[0]
                if index == 0
                else frame.add_paragraph()
            )

            paragraph.text = point
            paragraph.font.name = self.FONT
            paragraph.font.size = Pt(18)
            paragraph.font.color.rgb = self.GRAY_700

            paragraph.space_after = Pt(18)
            paragraph.level = 0

            paragraph.text = f"•  {point}"

        self._add_reference_footer(
            slide,
            slide_data,
        )

        self._add_notes(
            slide,
            slide_data.speaker_notes,
        )

    # =========================================================
    # STATISTIC SLIDE
    # =========================================================

    def _add_stat_slide(
        self,
        presentation: Presentation,
        slide_data: PresentationSlide,
    ) -> None:

        slide = self._create_base_slide(
            presentation,
            slide_data,
        )

        # Main content
        content_box = slide.shapes.add_textbox(
            Inches(0.85),
            Inches(1.85),
            Inches(7.1),
            Inches(4.5),
        )

        frame = content_box.text_frame
        frame.clear()
        frame.word_wrap = True

        for index, point in enumerate(slide_data.content):

            paragraph = (
                frame.paragraphs[0]
                if index == 0
                else frame.add_paragraph()
            )

            paragraph.text = f"•  {point}"
            paragraph.font.name = self.FONT
            paragraph.font.size = Pt(17)
            paragraph.font.color.rgb = self.GRAY_700
            paragraph.space_after = Pt(15)

        # Statistic card
        card = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Inches(8.55),
            Inches(2.0),
            Inches(3.8),
            Inches(3.2),
        )

        card.fill.solid()
        card.fill.fore_color.rgb = self.BLUE_LIGHT

        card.line.color.rgb = RGBColor(
            191,
            219,
            254,
        )

        label = slide.shapes.add_textbox(
            Inches(8.9),
            Inches(2.35),
            Inches(3.1),
            Inches(0.4),
        )

        self._set_text(
            label.text_frame,
            "KEY STATISTIC",
            font_size=10,
            bold=True,
            color=self.BLUE,
            align=PP_ALIGN.CENTER,
        )

        value = slide.shapes.add_textbox(
            Inches(8.85),
            Inches(2.95),
            Inches(3.2),
            Inches(1.3),
        )

        self._set_text(
            value.text_frame,
            slide_data.key_stat or "",
            font_size=25,
            bold=True,
            color=self.NAVY,
            align=PP_ALIGN.CENTER,
        )

        self._add_reference_footer(
            slide,
            slide_data,
        )

        self._add_notes(
            slide,
            slide_data.speaker_notes,
        )

    # =========================================================
    # TWO COLUMN SLIDE
    # =========================================================

    def _add_two_column_slide(
        self,
        presentation: Presentation,
        slide_data: PresentationSlide,
    ) -> None:

        slide = self._create_base_slide(
            presentation,
            slide_data,
        )

        midpoint = max(
            1,
            (len(slide_data.content) + 1) // 2,
        )

        left_points = slide_data.content[:midpoint]
        right_points = slide_data.content[midpoint:]

        self._add_column(
            slide,
            left_points,
            Inches(0.85),
            Inches(1.85),
            Inches(5.65),
        )

        if right_points:

            self._add_column(
                slide,
                right_points,
                Inches(6.85),
                Inches(1.85),
                Inches(5.65),
            )

        if slide_data.key_stat:

            self._add_small_stat(
                slide,
                slide_data.key_stat,
            )

        self._add_reference_footer(
            slide,
            slide_data,
        )

        self._add_notes(
            slide,
            slide_data.speaker_notes,
        )

    # =========================================================
    # TIMELINE SLIDE
    # =========================================================

    def _add_timeline_slide(
        self,
        presentation: Presentation,
        slide_data: PresentationSlide,
    ) -> None:

        slide = self._create_base_slide(
            presentation,
            slide_data,
        )

        count = len(slide_data.content)

        if count == 0:
            return

        start_x = 1.0
        end_x = 12.2

        y = 3.8

        # Timeline line
        line = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Inches(start_x),
            Inches(y),
            Inches(end_x - start_x),
            Inches(0.05),
        )

        line.fill.solid()
        line.fill.fore_color.rgb = self.GRAY_300
        line.line.fill.background()

        spacing = (
            (end_x - start_x) / max(count - 1, 1)
        )

        for index, point in enumerate(slide_data.content):

            x = (
                start_x
                if count == 1
                else start_x + spacing * index
            )

            marker = slide.shapes.add_shape(
                MSO_SHAPE.OVAL,
                Inches(x - 0.10),
                Inches(y - 0.10),
                Inches(0.20),
                Inches(0.20),
            )

            marker.fill.solid()
            marker.fill.fore_color.rgb = self.BLUE
            marker.line.fill.background()

            box_y = 2.25 if index % 2 == 0 else 4.15

            box = slide.shapes.add_textbox(
                Inches(max(0.55, x - 1.0)),
                Inches(box_y),
                Inches(2.0),
                Inches(1.25),
            )

            self._set_text(
                box.text_frame,
                point,
                font_size=12,
                bold=True,
                color=self.GRAY_700,
                align=PP_ALIGN.CENTER,
            )

        self._add_reference_footer(
            slide,
            slide_data,
        )

        self._add_notes(
            slide,
            slide_data.speaker_notes,
        )

    # =========================================================
    # COLUMN HELPER
    # =========================================================

    def _add_column(
        self,
        slide,
        points: list[str],
        x,
        y,
        width,
    ) -> None:

        card = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            x,
            y,
            width,
            Inches(4.3),
        )

        card.fill.solid()
        card.fill.fore_color.rgb = self.WHITE

        card.line.color.rgb = self.GRAY_200
        card.line.width = Pt(1)

        text_box = slide.shapes.add_textbox(
            x + Inches(0.3),
            y + Inches(0.35),
            width - Inches(0.6),
            Inches(3.6),
        )

        frame = text_box.text_frame
        frame.clear()
        frame.word_wrap = True

        for index, point in enumerate(points):

            paragraph = (
                frame.paragraphs[0]
                if index == 0
                else frame.add_paragraph()
            )

            paragraph.text = f"•  {point}"
            paragraph.font.name = self.FONT
            paragraph.font.size = Pt(15)
            paragraph.font.color.rgb = self.GRAY_700
            paragraph.space_after = Pt(14)

    # =========================================================
    # SMALL STAT
    # =========================================================

    def _add_small_stat(
        self,
        slide,
        value: str,
    ) -> None:

        card = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Inches(9.65),
            Inches(5.45),
            Inches(2.6),
            Inches(0.65),
        )

        card.fill.solid()
        card.fill.fore_color.rgb = self.BLUE_LIGHT
        card.line.fill.background()

        text = slide.shapes.add_textbox(
            Inches(9.8),
            Inches(5.57),
            Inches(2.3),
            Inches(0.35),
        )

        self._set_text(
            text.text_frame,
            value,
            font_size=11,
            bold=True,
            color=self.BLUE,
            align=PP_ALIGN.CENTER,
        )

    # =========================================================
    # BASE SLIDE
    # =========================================================

    def _create_base_slide(
        self,
        presentation: Presentation,
        slide_data: PresentationSlide,
    ):

        slide = presentation.slides.add_slide(
            presentation.slide_layouts[6]
        )

        background = slide.background.fill
        background.solid()
        background.fore_color.rgb = self.WHITE

        # Top accent
        accent = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            0,
            0,
            self.SLIDE_WIDTH,
            Inches(0.07),
        )

        accent.fill.solid()
        accent.fill.fore_color.rgb = self.BLUE
        accent.line.fill.background()

        # Title
        title = slide.shapes.add_textbox(
            Inches(0.85),
            Inches(0.55),
            Inches(10.5),
            Inches(0.7),
        )

        self._set_text(
            title.text_frame,
            slide_data.title,
            font_size=27,
            bold=True,
            color=self.NAVY,
        )

        # Subtitle
        if slide_data.subtitle:

            subtitle = slide.shapes.add_textbox(
                Inches(0.88),
                Inches(1.25),
                Inches(10.7),
                Inches(0.45),
            )

            self._set_text(
                subtitle.text_frame,
                slide_data.subtitle,
                font_size=12,
                color=self.GRAY_500,
            )

        self._add_slide_number(
            slide,
            slide_data.slide_number,
        )

        # Bottom footer
        footer = slide.shapes.add_textbox(
            Inches(0.85),
            Inches(6.9),
            Inches(7.0),
            Inches(0.25),
        )

        self._set_text(
            footer.text_frame,
            "PRISM",
            font_size=9,
            bold=True,
            color=self.GRAY_500,
        )

        return slide

    # =========================================================
    # SLIDE NUMBER
    # =========================================================

    def _add_slide_number(
        self,
        slide,
        slide_number: int,
    ) -> None:

        number = slide.shapes.add_textbox(
            Inches(11.75),
            Inches(0.52),
            Inches(0.75),
            Inches(0.3),
        )

        self._set_text(
            number.text_frame,
            f"{slide_number:02d}",
            font_size=10,
            bold=True,
            color=self.GRAY_500,
            align=PP_ALIGN.RIGHT,
        )

    # =========================================================
    # SOURCE REFERENCES
    # =========================================================

    def _add_reference_footer(
        self,
        slide,
        slide_data: PresentationSlide,
    ) -> None:

        references = getattr(
            slide_data,
            "source_references",
            [],
        )

        if not references:
            return

        source_text = "Sources: " + ", ".join(
            references[:3]
        )

        source_box = slide.shapes.add_textbox(
            Inches(7.4),
            Inches(6.82),
            Inches(5.0),
            Inches(0.3),
        )

        self._set_text(
            source_box.text_frame,
            source_text,
            font_size=8,
            color=self.GRAY_500,
            align=PP_ALIGN.RIGHT,
        )

    # =========================================================
    # SPEAKER NOTES
    # =========================================================

    def _add_notes(
        self,
        slide,
        notes: str,
    ) -> None:

        if not notes:
            return

        notes_slide = slide.notes_slide

        notes_frame = notes_slide.notes_text_frame

        notes_frame.text = notes

    # =========================================================
    # TEXT HELPER
    # =========================================================

    def _set_text(
        self,
        text_frame,
        text: str,
        *,
        font_size: int,
        color,
        bold: bool = False,
        align=PP_ALIGN.LEFT,
    ) -> None:

        text_frame.clear()

        text_frame.word_wrap = True

        text_frame.vertical_anchor = MSO_ANCHOR.TOP

        text_frame.margin_left = 0
        text_frame.margin_right = 0
        text_frame.margin_top = 0
        text_frame.margin_bottom = 0

        paragraph = text_frame.paragraphs[0]

        paragraph.text = text
        paragraph.alignment = align

        paragraph.font.name = self.FONT
        paragraph.font.size = Pt(font_size)
        paragraph.font.bold = bold
        paragraph.font.color.rgb = color